#!/usr/bin/env python3
"""Standalone CET XLSX/API quality test. Python >=3.10; openpyxl==3.1.5."""
import argparse
import glob
import gzip
import importlib.metadata
import os
from pathlib import Path
import platform
import sys
import time
import threading
import uuid
import urllib.error
import urllib.parse
import urllib.request

import collections
import datetime as dt
import hashlib
import json
import math
from zoneinfo import ZoneInfo

UTC = dt.timezone.utc
STEP = dt.timedelta(minutes=15)

def iso(value):
    return value.astimezone(UTC).isoformat(timespec='seconds').replace('+00:00', 'Z')


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n')


def excel_datetime(value):
    from openpyxl.utils.datetime import from_excel
    if isinstance(value, (float, int)):
        value = from_excel(value)
    if not isinstance(value, dt.datetime) or value.tzinfo is not None:
        raise ValueError('Expected a naive Excel datetime')
    rounded = value.replace(microsecond=0)
    if value.microsecond >= 999000:
        rounded += dt.timedelta(seconds=1)
    elif value.microsecond > 1000:
        raise ValueError('Timestamp has material fractional seconds')
    if rounded.second or rounded.minute % 15:
        raise ValueError('Timestamp is not aligned to PT15M')
    return rounded


def read_workbook(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Exactly one worksheet expected')
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        columns = {name: headers.index(name) for name in
                   ['Meldepunkt', 'OBIS', 'Datum von', 'Datum bis', 'Wert']}
        records, errors = [], []
        skipped = 0
        meters, obis = set(), set()
        for number, row in enumerate(rows, 2):
            if all(v is None for v in row):
                skipped += 1
                continue
            try:
                data = {k: row[i] if i < len(row) else None for k, i in columns.items()}
                start, end = excel_datetime(data['Datum von']), excel_datetime(data['Datum bis'])
                value = data['Wert']
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                    raise ValueError('Missing/non-numeric Wert; not replaced by zero')
                if data['Meldepunkt'] is None or data['OBIS'] is None:
                    raise ValueError('Missing meter or OBIS')
                meters.add(str(data['Meldepunkt']))
                obis.add(str(data['OBIS']))
                records.append((start, end, float(value), number))
            except (ValueError, TypeError, IndexError) as error:
                if len(errors) < 20:
                    errors.append({'row': number, 'error': str(error)})
        if errors:
            raise ValueError(f'Invalid input rows (first 20): {errors}')
        if not records or len(meters) != 1 or len(obis) != 1:
            raise ValueError('One nonempty meter/OBIS series per workbook required')
        counts = collections.Counter(r[0] for r in records)
        audit = {
            'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'sheet': sheet.title, 'records': len(records), 'blank_rows_ignored': skipped,
            'meter': next(iter(meters)), 'obis': next(iter(obis)),
            'first_wall_time': min(counts).isoformat(), 'last_wall_time': max(counts).isoformat(),
            'duplicate_wall_times': sum(n - 1 for n in counts.values()),
            'non_15_minute_wall_durations': sum(e - s != STEP for s, e, _, _ in records),
        }
        return records, audit
    finally:
        workbook.close()


def local_candidates(wall, zone):
    candidates = set()
    for fold in (0, 1):
        candidate = wall.replace(tzinfo=zone, fold=fold).astimezone(UTC)
        if candidate.astimezone(zone).replace(tzinfo=None) == wall:
            candidates.add(candidate)
    return sorted(candidates)


def dataset_from_records(records, audit, unit, basis):
    zone = ZoneInfo({'berlin': 'Europe/Berlin', 'cet': 'Etc/GMT-1', 'utc': 'UTC'}[basis])
    counts = collections.Counter(r[0] for r in records)
    occurrences = collections.Counter()
    values, seen = [], set()
    previous = None
    for start, end, value, row in records:
        candidates = local_candidates(start, zone)
        if not candidates:
            raise ValueError(f'Row {row}: nonexistent local time {start}; clarify export time convention')
        if len(candidates) == 2:
            if counts[start] != 2:
                raise ValueError(f'Row {row}: ambiguous autumn time {start} without two occurrences')
            instant = candidates[occurrences[start]]
            occurrences[start] += 1
        else:
            instant = candidates[0]
        if instant in seen:
            raise ValueError(f'Row {row}: duplicate physical interval; no deduplication performed')
        if previous is not None and instant <= previous:
            raise ValueError(f'Row {row}: input is not chronologically ordered')
        if instant + STEP not in local_candidates(end, zone):
            raise ValueError(f'Row {row}: Datum bis is not the end of a physical PT15M interval')
        seen.add(instant)
        previous = instant
        values.append({'timestamp': iso(instant), 'value': value})
    audit.update(unit=unit, time_basis=basis, timezone=zone.key,
                 interval_semantics='Datum von=start; Datum bis=exclusive end')
    return {'series_id': 'meter-' + audit['meter'], 'unit': unit, 'timezone': zone.key,
            'period_from': records[0][0].date().isoformat(),
            'period_until': records[-1][0].date().isoformat(), 'values': values}



def verify_and_score(result, dataset, start, end):
    """Independently verify API outputs and compute pooled sufficient statistics."""
    zone = ZoneInfo(dataset['timezone'])
    expected = {}
    cursor = dt.datetime.combine(start, dt.time(), zone).astimezone(UTC)
    stop = dt.datetime.combine(end + dt.timedelta(days=1), dt.time(), zone).astimezone(UTC)
    while cursor < stop:
        expected[iso(cursor)] = cursor
        cursor += STEP
    actuals = {r['timestamp']: r['value'] for r in dataset['values']}
    squared = absolute = signed = actual_abs = percentage = 0.0
    count = percentage_count = 0
    observed = set()
    for row in result['forecast_values']:
        stamp = iso(dt.datetime.fromisoformat(row['timestamp'].replace('Z', '+00:00')))
        if stamp not in expected or stamp in observed:
            raise ValueError('API returned an unexpected or duplicate forecast interval')
        observed.add(stamp)
        day = expected[stamp].astimezone(zone).date()
        cutoff = dt.datetime.combine(day - dt.timedelta(days=1), dt.time(), zone).astimezone(UTC)
        trained = dt.datetime.fromisoformat(row['training_data_until'].replace('Z', '+00:00'))
        issued = dt.datetime.fromisoformat(row['forecast_created_at'].replace('Z', '+00:00'))
        issue_time = result.get('forecast_run', {}).get('issue_time', '00:00')
        if issue_time not in ('00:00', '07:00', '18:00'):
            raise ValueError('Unsupported issue_time')
        expected_issue = dt.datetime.combine(day-dt.timedelta(days=1), dt.time.fromisoformat(issue_time), zone).astimezone(UTC)
        if trained >= cutoff or issued != expected_issue:
            raise ValueError('API violates D-2 training/issuance contract')
        forecast = row['predicted_value']
        if not isinstance(forecast, (int, float)) or not math.isfinite(forecast):
            raise ValueError('API returned a non-finite prediction')
        actual = actuals.get(stamp)
        if row.get('actual_value') != actual:
            raise ValueError('API actual differs from original imported measurement')
        if actual is None:
            continue
        error = actual - forecast
        count += 1
        squared += error * error
        absolute += abs(error)
        signed += error
        actual_abs += abs(actual)
        if actual != 0:
            percentage += abs(error / actual) * 100
            percentage_count += 1
    if len(observed) != len(expected) or not count:
        raise ValueError('Forecast incomplete or no holdout actuals matched')
    scores = {'sample_count': count, 'expected_intervals': len(expected), 'coverage': count / len(expected),
              'mse': squared / count, 'rmse': math.sqrt(squared / count), 'mae': absolute / count,
              'bias': signed / count, 'cumulative_error': signed,
              'mape': percentage / percentage_count if percentage_count else None,
              'wape_percent': 100 * absolute / actual_abs if actual_abs else (0 if not absolute else None),
              'squared_error_sum': squared, 'absolute_error_sum': absolute, 'actual_absolute_sum': actual_abs,
              'mape_sample_count': percentage_count}
    for metric in ('mse', 'rmse', 'mae', 'bias', 'cumulative_error'):
        if not math.isclose(scores[metric], result['backtest'][metric], abs_tol=1e-8, rel_tol=1e-8):
            raise ValueError(f'API metric mismatch: {metric}')
    return scores


VERSION = '3.1.0'
API_PATH = '/api/forecast-sandbox/consumption/evaluation'
DEFAULT_CONFIG = {
    'mode': 'rolling_day_ahead', 'relationship_mode': 'auto',
    'model_family': 'learned_states', 'extended_features': True,
    'feature_set': ['history', 'calendar'], 'issue_time': '00:00',
    'selection_metric': 'mae', 'recent_window_days': 84,
}
OVERRIDES = {'model_family', 'extended_features', 'feature_set', 'issue_time',
             'selection_metric', 'recent_window_days', 'weather_dataset_id', 'context_dataset_id',
             'selection_policy', 'activity_model', 'activity_labeling', 'prediction_threshold_w'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def encoded(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False,
                      sort_keys=True, separators=(',', ':')).encode('utf-8')


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2) + '\n', encoding='utf-8')


class RejectRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        raise ValueError('HTTP redirect rejected; use the final API URL')


class ApiArchive:
    def __init__(self, base, out, timeout, replay=None):
        self.base, self.out, self.timeout, self.replay = base.rstrip('/'), out, timeout, replay
        self.calls = []
        self.archived = {}
        if replay:
            self.archived = {c['name']: c for c in json.loads((replay / 'http-calls.json').read_text())}
        url = urllib.parse.urlsplit(self.base)
        if (url.scheme not in ('http', 'https') or not url.hostname or url.username or url.password
                or url.path or url.query or url.fragment):
            raise ValueError('Expected HTTP(S) base URL without path or credentials')
        if url.scheme == 'http' and url.hostname not in ('localhost', '127.0.0.1', '::1'):
            raise ValueError('Use HTTPS for remote APIs')
        self.opener = urllib.request.build_opener(RejectRedirect())

    def call(self, name, endpoint, payload=None, accepted_statuses=(200,)):
        body = encoded(payload) if payload is not None else None
        if body and len(body) > 25 * 1024 * 1024:
            raise ValueError('Request exceeds 25 MiB; input is not truncated')
        request_hash = sha(body) if body else None
        started = time.monotonic()
        if self.replay:
            entry = self.archived[name]
            if entry['path'] != endpoint or entry['request_sha256'] != request_hash:
                raise ValueError('Replay request differs from original test')
            with gzip.open(self.replay / (name + '.response.json.gz'), 'rb') as f:
                raw = f.read()
            if sha(raw) != entry['response_sha256']:
                raise ValueError('Archived API response checksum mismatch')
            status = entry['status']
        else:
            headers = {'Content-Type': 'application/json', 'Accept': 'application/json'}
            if os.environ.get('CET_API_TOKEN'):
                headers['Authorization'] = 'Bearer ' + os.environ['CET_API_TOKEN']
            request = urllib.request.Request(self.base + endpoint, data=body, headers=headers)
            try:
                response = self.opener.open(request, timeout=self.timeout)
            except urllib.error.HTTPError as error:
                response = error
            with response:
                raw, status = response.read(), response.code
        if body:
            with gzip.open(self.out / (name + '.request.json.gz'), 'wb') as f:
                f.write(body)
        with gzip.open(self.out / (name + '.response.json.gz'), 'wb') as f:
            f.write(raw)
        self.calls.append({'name': name, 'path': endpoint, 'status': status,
                           'request_sha256': request_hash, 'response_sha256': sha(raw),
                           'elapsed_seconds': round(time.monotonic() - started, 3)})
        save(self.out / 'http-calls.json', self.calls)
        if status not in accepted_statuses:
            raise ValueError(f'API returned HTTP {status}; see {name}.response.json.gz')
        return json.loads(raw)


def cell(value):
    return str(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('|', '&#124;').replace('\n', ' ').replace('\r', ' ')


READINESS_POLICY = {
    'version': 'report_readiness_v1',
    'meaning': 'Descriptive backtest heuristic, not model skill versus a baseline or production readiness',
    'strong_max_wape_percent': 10,
    'medium_max_wape_percent': 30,
    'minimum_days_for_strong': 28,
    'low_load_mean_abs_kwh_per_interval': 0.1,
    'small_mae_kwh_per_interval': 0.1,
    'small_rmse_kwh_per_interval': 0.15,
}


def interpret_readiness(metrics, result, unit, start, until):
    """Describe observed errors using fixed, disclosed thresholds; never change forecasts."""
    count = metrics['sample_count']
    mean_abs = metrics['actual_absolute_sum'] / count
    days = (until - start).days + 1
    factor = 0.25 if unit == 'kW' else 1.0
    low_load = mean_abs * factor <= READINESS_POLICY['low_load_mean_abs_kwh_per_interval']
    small_errors = (metrics['mae'] * factor <= READINESS_POLICY['small_mae_kwh_per_interval']
                    and metrics['rmse'] * factor <= READINESS_POLICY['small_rmse_kwh_per_interval'])
    wape = metrics['wape_percent']
    signal = ('weak' if wape is None or wape > READINESS_POLICY['medium_max_wape_percent']
              else 'strong' if wape <= READINESS_POLICY['strong_max_wape_percent']
              else 'medium')
    if days < READINESS_POLICY['minimum_days_for_strong'] and signal == 'strong':
        signal = 'medium'
    if wape is None:
        interpretation = 'Keine messbare Lastbasis im Prüfzeitraum; WAPE und relatives Forecast-Signal nicht beurteilbar.'
        step = 'Zeitraum mit realen Lastaktivitäten prüfen; Messung und Stillstandszeiten fachlich einordnen.'
    elif low_load and small_errors and wape > READINESS_POLICY['medium_max_wape_percent']:
        interpretation = (f'Kleine absolute Fehler bei geringer Lastbasis (mittlerer |Istwert| {mean_abs:.6f} {unit}/Intervall); '
                          'der Fehler ist relativ zur Last dennoch hoch. Die Prozentmetrik allein reicht zur betrieblichen Bewertung nicht aus.')
        step = 'Absolute Fehlertoleranz festlegen; Null-/Aktivintervalle und Tagesaggregation getrennt prüfen und mit einer einfachen Baseline vergleichen.'
    elif signal == 'weak':
        interpretation = 'Hoher relativer Fehler im geprüften Zeitraum; eingeschränktes Forecast-Signal nach der Berichtheuristik.'
        step = 'Fehlertage und Lastwechsel analysieren; gegen eine einfache Baseline prüfen und Verbesserungen zeitlich getrennt backtesten.'
    elif signal == 'medium':
        interpretation = 'Mittleres Forecast-Signal nach der Berichtheuristik; Aussage gilt nur für den geprüften Zeitraum.'
        step = 'Fehler nach Tageszeit und Lastzustand prüfen; Baseline-Vergleich und längeren rollierenden Test ergänzen.'
    else:
        interpretation = 'Niedriger relativer Fehler im geprüften Zeitraum; starkes Signal nach der Berichtheuristik, keine Produktivfreigabe.'
        step = 'Gegen eine einfache Baseline bestätigen und betriebliche Fehlertoleranzen sowie saisonale Stabilität prüfen.'
    if days < READINESS_POLICY['minimum_days_for_strong']:
        interpretation += f' Kurzer Prüfzeitraum ({days} Tage); keine belastbare Aussage zur langfristigen Stabilität.'
        step += ' Auf mindestens 28 Tage und anschließend saisonal repräsentative Zeiträume erweitern.'
    zero_fraction = sum(v.get('actual_value') == 0 for v in result['forecast_values']) / count
    return {'forecast_signal': signal, 'interpretation': interpretation, 'recommended_next_step': step,
            'readiness_evidence': {'policy_version': READINESS_POLICY['version'],
                                   'mean_absolute_actual': mean_abs, 'unit': unit,
                                   'zero_interval_fraction': zero_fraction, 'evaluation_days': days,
                                   'low_load': low_load, 'small_absolute_errors': small_errors,
                                   'baseline_comparison_available': False}}


def archive_run_inputs(out, script_bytes, config, input_bytes=None):
    """Archive captured inputs; hashes describe the exact written bytes."""
    artifacts = {'xlsx_forecast_test.py': script_bytes,
                 'configuration.json': (json.dumps(config, ensure_ascii=False, allow_nan=False,
                                                   indent=2) + '\n').encode('utf-8')}
    if input_bytes is not None:
        artifacts['configuration.input.json'] = input_bytes
    for name, data in artifacts.items():
        (out / name).write_bytes(data)
    return {name: {'sha256': sha(data), 'bytes': len(data)} for name, data in artifacts.items()}


def feature_usage(result):
    values = result['forecast_values']
    counts = collections.Counter()
    predictors = collections.Counter()
    unknown = 0
    for value in values:
        if 'feature_set' not in value:
            unknown += 1
        counts.update(set(value.get('feature_set', [])))
        predictors.update(set(value.get('selected_predictors', [])))
    return {'interval_count': len(values), 'unknown_intervals': unknown,
            'feature_interval_counts': dict(sorted(counts.items())),
            'predictor_interval_counts': dict(sorted(predictors.items()))}


def markdown(report):
    lines = ['# Forecast-Readiness-Backtest – XLSX/API-Test', '',
             f"Aufruf (UTC): {report['invoked_at']} · Status: **{report['status']}**",
             f"Script SHA-256: `{report['script_sha256']}`",
             f"Testzeitraum: {report['from']} bis {report['until']} (einschließlich) · Einheit: {report['unit']} · Zeitbasis: {report['time_basis']}",
             f"Ausführung: {report['execution']} · API: {cell(report['base_url'])}", '',
             f"Wetterregion: `{cell(report.get('weather_region', 'nicht dokumentiert'))}`",
             f"Angefordertes feature_set: `{', '.join(report.get('configuration', {}).get('feature_set', []))}`",
             '', 'Effektive Modellkonfiguration (einschließlich Client-Defaults):', '', '```json',
             json.dumps(report.get('configuration', {}), ensure_ascii=False, indent=2), '```', '',
             '| Quelldatei (XLSX) | RMSE | MAE | WAPE % | Status | forecast_signal | Genutzte Features (% Intervalle) |',
             '|---|---:|---:|---:|---|---|---|']
    for row in report['files']:
        metrics = row.get('metrics', {})
        numbers = ['—' if metrics.get(key) is None else f'{metrics[key]:.6f}'
                   for key in ('rmse', 'mae', 'wape_percent')]
        usage = row.get('feature_usage', {})
        total = usage.get('interval_count', 0)
        used = ', '.join(f'{name} {100 * count / total:.1f}%' for name, count in
                         usage.get('feature_interval_counts', {}).items()) if total else ''
        if usage.get('unknown_intervals'):
            used += f"; unbekannt: {usage['unknown_intervals']} Intervalle"
        lines.append('| ' + cell(row['file']) + ' | ' + ' | '.join(numbers) + ' | ' + cell(row['status']) + ' | ' + {'weak': '🔴 weak', 'medium': '🟡 medium', 'strong': '🟢 strong'}.get(row.get('forecast_signal'), '—') + ' | ' + cell(used or '—') + ' |')
    lines += ['', 'Genutzte Features laut API je Prognoseintervall; angeforderte Kandidaten werden nur bei Auswahl und Verfügbarkeit verwendet. Detail-Prädiktoren und Quellen stehen in report.json.']
    comparisons = [row for row in report['files'] if row.get('filter_comparison')]
    if comparisons:
        lines += ['', 'Schwellwertvergleich: Haupttabelle und Ampel bewerten die gefilterte Prognose. Istwerte bleiben unverändert. Werte genau an der Grenze bleiben erhalten.', '',
                  '| XLSX | Grenze W | Auf null gesetzte Intervalle | RMSE roh | RMSE gefiltert | MAE roh | MAE gefiltert | WAPE % roh | WAPE % gefiltert |',
                  '|---|---:|---:|---:|---:|---:|---:|---:|---:|']
        for row in comparisons:
            comparison = row['filter_comparison']
            raw, filtered = comparison['raw_metrics'], row['metrics']
            values = [metrics.get(key) for key in ('rmse', 'mae', 'wape_percent') for metrics in (raw, filtered)]
            numbers = ['—' if value is None else f'{value:.6f}' for value in values]
            lines.append(f"| {cell(row['file'])} | {comparison['threshold_w']} | {comparison['changed_intervals']} | " + ' | '.join(numbers) + ' |')
    for name, artifact in report.get('artifacts', {}).items():
        lines += ['', f"Archiv `{name}` · SHA-256: `{artifact['sha256']}`"]
    for row in report['files']:
        if row.get('interpretation'):
            lines += ['', f"**{cell(row['file'])}** — interpretation: {cell(row['interpretation'])}",
                      f"recommended_next_step: {cell(row['recommended_next_step'])}"]
    lines += ['', 'Ampel = Berichtheuristik: WAPE ≤10 % strong, ≤30 % medium, sonst weak; unter 28 Prüftagen höchstens medium. Ohne Lastbasis weak = nicht beurteilbar. Kein Leistungsversprechen und kein Nachweis eines Vorteils gegenüber einer Baseline.',
              'Geringe Lastbasis: mittlerer |Istwert| ≤0,1 kWh/Intervall; kleine absolute Fehler: MAE ≤0,1 und RMSE ≤0,15 kWh (kW äquivalent ×0,25 h). Feste Orientierungswerte, keine betriebliche Fehlertoleranz.',
              'RMSE/MAE in der Eingabeeinheit je Viertelstunde. WAPE = 100 × Σ|Ist−Prognose| / Σ|Ist|; bei Σ|Ist|=0 nicht definiert.',
              'Rollierender Backtest, Zählerhistorie bis D−2; kein Trainingsfehler. Kennzahlen unabhängig aus den API-Werten nachgerechnet.',
              'Datei-Prüfsummen, Konfiguration, Modellversionen und Warnungen: report.json; vollständige Anfragen/Antworten: *.json.gz.',
              'Identische Live-Ergebnisse erfordern dieselben XLSX, Script-, Server- und Featuredaten-Versionen. Replay prüft archivierte Ergebnisse ohne neuen API-Lauf.']
    for row in report['files']:
        if row.get('error'):
            lines.append(f"Fehler {cell(row['file'])}: {cell(row['error'])}")
    if report.get('error'):
        lines.append('Fehler: ' + cell(report['error']))
    return '\n'.join(lines) + '\n'


def publish_report(out, report):
    save(out / 'report.json', report)
    (out / 'report.md').write_text(markdown(report), encoding='utf-8')


def expand_files(patterns):
    files = []
    for pattern in patterns:
        matches = sorted(glob.glob(pattern)) if glob.has_magic(pattern) else [pattern]
        if not matches:
            raise ValueError(f'No files match: {pattern}')
        for item in matches:
            path = Path(item).resolve()
            if path.suffix.lower() != '.xlsx' or not path.is_file():
                raise ValueError(f'Expected an existing XLSX file: {item}')
            if path not in files:
                files.append(path)
    return files


def verify_threshold(result, dataset, config, start, until):
    watts = config.get('prediction_threshold_w', 0)
    if result['forecast_run'].get('prediction_threshold_w', 0) != watts:
        raise ValueError('API prediction threshold differs from requested configuration')
    if not watts:
        return None
    threshold = watts / 1000 * (0.25 if dataset['unit'] == 'kWh' else 1)
    changed = 0
    raw_values = []
    for row in result['forecast_values']:
        raw = row.get('raw_predicted_value')
        if type(raw) not in (int, float) or not math.isfinite(raw):
            raise ValueError('Missing/non-finite raw prediction for threshold audit')
        zeroed = raw != 0 and abs(raw) < threshold
        expected = 0 if zeroed else raw
        if row['predicted_value'] != expected or row.get('prediction_zeroed') is not zeroed:
            raise ValueError('API threshold mapping is incorrect')
        changed += int(zeroed)
        raw_values.append({**row, 'predicted_value': raw})
    audit = result.get('filter_evaluation', {})
    if (audit.get('changed_intervals') != changed or audit.get('actuals_modified') is not False
            or audit.get('threshold_w') != watts or audit.get('threshold_value') != threshold):
        raise ValueError('API threshold audit metadata is inconsistent')
    raw_result = {**result, 'forecast_values': raw_values, 'backtest': audit['raw_backtest']}
    raw_metrics = verify_and_score(raw_result, dataset, start, until)
    if raw_metrics['actual_absolute_sum'] == 0:
        raw_metrics['wape_percent'] = None
    return {'threshold_w': watts, 'threshold_value': threshold, 'unit': dataset['unit'],
            'changed_intervals': changed, 'actuals_modified': False, 'raw_metrics': raw_metrics}


def score_response(response, dataset, config, start, until):
    results = response.get('results', [])
    if len(results) != 1:
        raise ValueError('Expected one API result per source file')
    result = results[0]
    run = result['forecast_run']
    if config.get('selection_policy') == 'adaptive_rmse_v1':
        if run.get('model_version') != 'relationship_state_adaptive_v5':
            raise ValueError('API did not execute adaptive_rmse_v1; compatible v5 deployment required')
        selections = result.get('relationship_analysis', {}).get('selections', [])
        if not selections or any('adaptive_guard' not in s.get('reference_selection', s) for s in selections):
            raise ValueError('API omitted adaptive incumbent-guard evidence')
    for key, value in {'series_id': dataset['series_id'], 'timezone': dataset['timezone'],
                       'mode': 'rolling_day_ahead', 'relationship_mode': 'auto',
                       'forecast_period_from': start.isoformat(), 'forecast_period_until': until.isoformat(),
                       'model_family': config['model_family'], 'issue_time': config['issue_time']}.items():
        if run.get(key, '00:00' if key == 'issue_time' else None) != value:
            raise ValueError(f'Unexpected API metadata: {key}')
    for value in result['forecast_values']:
        if value.get('series_id') != dataset['series_id'] or value.get('unit') != dataset['unit']:
            raise ValueError('Forecast series/unit mismatch')
        stamp = dt.datetime.fromisoformat(value['timestamp'].replace('Z', '+00:00'))
        if value.get('forecast_for') != stamp.astimezone(ZoneInfo(dataset['timezone'])).date().isoformat():
            raise ValueError('Forecast date/timestamp mismatch')
    scores = verify_and_score(result, dataset, start, until)
    if scores['coverage'] != 1:
        raise ValueError('Missing actual values in test period; full coverage required')
    if scores['actual_absolute_sum'] == 0:
        scores['wape_percent'] = None
    comparison = verify_threshold(result, dataset, config, start, until)
    if comparison is not None:
        scores['filter_comparison'] = comparison
    return scores, result


def run_forecast_job(api, name, payload, progress, on_job, poll_interval=5, job_timeout=86400, run_path=None):
    response = progress.run('Forecast-Job einreichen', lambda: api.call(
        name, run_path or API_PATH + '/run', payload, accepted_statuses=(200, 202)), api_wait=not bool(api.replay))
    if 'jobId' not in response:
        # Old synchronous APIs and archives remain verifiable.
        if api.calls[-1]['status'] != 200:
            raise ValueError('HTTP 202 response lacks jobId')
        return response
    job_id = str(uuid.UUID(response['jobId']))
    status_url = f'/api/jobs/{job_id}/status'
    result_url = f'/api/jobs/{job_id}/result'
    cancel_url = API_PATH + f'/jobs/{job_id}/cancel'
    for key, expected in [('statusUrl', status_url), ('resultUrl', result_url),
                          ('progressUrl', f'/api/jobs/{job_id}/progress')]:
        if response.get(key) != expected:
            raise ValueError(f'Unexpected job URL: {key}')
    on_job({**response, 'cancelUrl': cancel_url})
    progress.emit(f'Job {job_id} angenommen; Status: {status_url}')
    started = time.monotonic()
    count = 0
    try:
        while True:
            if time.monotonic() - started > job_timeout:
                raise TimeoutError(f'Job {job_id}: client wait limit exceeded')
            count += 1
            status = api.call(f'{name}-status-{count:05d}', status_url)
            if status.get('jobId') != job_id:
                raise ValueError('Job status ID mismatch')
            state = status.get('status')
            details = status.get('progress') or {}
            step, total = details.get('step'), details.get('totalSteps')
            completed = f'{step}/{total} Prognosetage abgeschlossen' if step is not None and total is not None else 'Tagesfortschritt noch unbekannt'
            age = details.get('at') or status.get('updatedAt') or 'unbekannt'
            progress.emit(f"Job {state}: {completed}; {details.get('message') or status.get('phase') or 'wartet'}; letzter Fortschritt {age}")
            if state == 'completed':
                return progress.run('Job-Ergebnis laden', api.call, name + '-result', result_url, api_wait=not bool(api.replay))
            if state in ('error', 'failed', 'cancelled', 'recovery_pending'):
                raise ValueError(f"Job {job_id}: {state}: {status.get('error') or status.get('recovery') or 'see job status'}")
            if state not in ('queued', 'running'):
                raise ValueError(f'Unknown job state: {state}')
            if not api.replay:
                time.sleep(poll_interval)
    except (KeyboardInterrupt, TimeoutError):
        if not api.replay:
            try:
                api.call(name + '-cancel', cancel_url, {})
                progress.emit(f'Abbruch für Job {job_id} angefordert')
            except Exception as error:
                progress.emit(f'Job-Abbruch nicht bestätigt: {error}; POST {cancel_url} zum manuellen Abbruch')
        raise


class CliProgress:
    """Client-side stage progress; does not infer server progress from elapsed time."""
    def __init__(self, total, interval=15, stream=None):
        self.total = total
        self.interval = interval
        self.stream = stream if stream is not None else sys.stderr
        self.started = time.monotonic()
        self.index = 0
        self.file = 'Initialisierung'
        self.completed = 0
        self.failed = 0

    @staticmethod
    def duration(seconds):
        seconds = max(0, int(seconds))
        return f'{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}'

    def emit(self, message):
        name = self.file.replace('\n', ' ').replace('\r', ' ')
        print(f'[{self.index}/{self.total}] {name} | {message} | '
              f'Gesamt {self.duration(time.monotonic() - self.started)} | '
              f'{self.completed} erfolgreich, {self.failed} fehlgeschlagen',
              file=self.stream, flush=True)

    def run(self, label, action, *args, api_wait=False):
        started = time.monotonic()
        stop = threading.Event()
        def heartbeat():
            while not stop.wait(self.interval):
                suffix = 'warte auf API-Antwort; Serverfortschritt unbekannt' if api_wait else 'in Bearbeitung'
                self.emit(f'{label}: {suffix}, Schritt {self.duration(time.monotonic() - started)}')
        self.emit(label + ': gestartet')
        worker = None
        if self.interval > 0:
            worker = threading.Thread(target=heartbeat, daemon=True, name='xlsx-cli-progress')
            worker.start()
        outcome = 'abgebrochen/fehlgeschlagen'
        try:
            result = action(*args)
            outcome = 'fertig'
            return result
        finally:
            stop.set()
            if worker is not None:
                worker.join()
            self.emit(f'{label}: {outcome}, Schritt {self.duration(time.monotonic() - started)}')


def main(argv=None):
    invoked = iso(dt.datetime.now(UTC))
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--xlsx', nargs='+', required=True, help='One or more XLSX paths or quoted glob patterns')
    parser.add_argument('--from', dest='start', type=dt.date.fromisoformat, default=dt.date(2024, 1, 1))
    parser.add_argument('--until', type=dt.date.fromisoformat, default=dt.date(2024, 12, 31))
    parser.add_argument('--unit', choices=['kWh', 'kW'], default='kWh')
    parser.add_argument('--time-basis', choices=['berlin', 'cet', 'utc'], default='berlin')
    parser.add_argument('--base-url', default=os.environ.get('CET_BASE_URL', 'http://localhost:3900'))
    parser.add_argument('--config', type=Path, help='Optional JSON feature/model overrides')
    parser.add_argument('--weather-region', default='DE-BY-Kempten-87435')
    parser.add_argument('--server-revision', default='not_exposed_by_API', help='Optional operator-supplied server build identifier')
    parser.add_argument('--timeout', type=int, default=60, help='HTTP socket timeout per request (seconds)')
    parser.add_argument('--poll-interval', type=int, default=5, help='Job polling interval in seconds')
    parser.add_argument('--job-timeout', type=int, default=86400, help='Maximum client wait for one job (seconds); requests cancellation on expiry')
    parser.add_argument('--progress-interval', type=int, default=15,
                        help='Seconds between stderr status messages; 0 disables periodic messages')
    parser.add_argument('--out', type=Path, required=True, help='New output directory; never overwritten')
    parser.add_argument('--replay', type=Path, help='Verify a previous archive; no HTTP calls')
    args = parser.parse_args(argv)
    try:
        files = expand_files(args.xlsx)
        config = dict(DEFAULT_CONFIG)
        input_bytes = None
        if args.config:
            input_bytes = args.config.read_bytes()
            overrides = json.loads(input_bytes.decode('utf-8'))
            if not isinstance(overrides, dict) or set(overrides) - (OVERRIDES | {'mode', 'relationship_mode'}):
                raise ValueError('Unsupported configuration fields')
            config.update(overrides)
            if config['mode'] != 'rolling_day_ahead' or config['relationship_mode'] != 'auto':
                raise ValueError('Only rolling_day_ahead / auto supported by this test')
        original = None
        if args.replay:
            if args.config:
                raise ValueError('Replay uses original configuration; --config is not allowed')
            original = json.loads((args.replay / 'report.json').read_text(encoding='utf-8'))
            if original['status'] != 'completed':
                raise ValueError('Replay requires a completed archive')
            config = original['configuration']
            args.start, args.until = dt.date.fromisoformat(original['from']), dt.date.fromisoformat(original['until'])
            args.unit, args.time_basis = original['unit'], original['time_basis']
            args.weather_region = original['weather_region']
            args.base_url = original['base_url']
            if [sha(f.read_bytes()) for f in files] != [f['sha256'] for f in original['files']]:
                raise ValueError('Replay input checksums/order differ from original')
        if args.poll_interval <= 0 or args.job_timeout <= 0:
            raise ValueError('poll-interval and job-timeout must be positive')
        if args.progress_interval < 0:
            raise ValueError('progress-interval must be nonnegative')
        if args.until < args.start or args.timeout <= 0:
            raise ValueError('Invalid period/timeout')
        threshold = config.get('prediction_threshold_w', 0)
        if type(threshold) not in (int, float) or not math.isfinite(threshold) or threshold < 0:
            raise ValueError('prediction_threshold_w must be finite and nonnegative')
        if config['issue_time'] not in ('00:00', '07:00', '18:00'):
            raise ValueError('issue_time must be 00:00, 07:00 or 18:00')
        script_bytes = Path(__file__).read_bytes()
        args.out.mkdir(parents=True, exist_ok=False)
    except (ValueError, OSError, KeyError) as error:
        parser.error(str(error))
    report = {'version': VERSION, 'status': 'running', 'invoked_at': invoked,
              'execution': 'archived_api_replay' if original else 'live_http',
              'script_sha256': sha(script_bytes), 'base_url': args.base_url,
              'server_revision_operator_supplied': args.server_revision,
              'from': args.start.isoformat(), 'until': args.until.isoformat(),
              'unit': args.unit, 'time_basis': args.time_basis, 'weather_region': args.weather_region,
              'configuration': config, 'configuration_sha256': sha(encoded(config)),
              'readiness_policy': READINESS_POLICY,
              'execution_options': {'http_timeout_seconds': args.timeout, 'poll_interval_seconds': args.poll_interval,
                                    'job_timeout_seconds': args.job_timeout},
              'python': platform.python_version(),
              'dependencies': {name: importlib.metadata.version(name) for name in ('openpyxl', 'et-xmlfile')},
              'files': [{'file': f.name, 'sha256': sha(f.read_bytes()), 'status': 'pending'} for f in files]}
    if original:
        report['replay_origin'] = {'invoked_at': original['invoked_at'], 'script_sha256': original['script_sha256'],
                                   'report_sha256': sha((args.replay / 'report.json').read_bytes())}
    report['artifacts'] = archive_run_inputs(args.out, script_bytes, config, input_bytes)
    publish_report(args.out, report)
    progress = CliProgress(len(files), args.progress_interval)
    progress.emit(f'API-Socket-Timeout: {args.timeout}s; keine automatische POST-Wiederholung')
    try:
        api = ApiArchive(args.base_url, args.out, args.timeout, args.replay)
        spec = progress.run('API prüfen', api.call, 'openapi', '/api/forecast-sandbox/openapi.json', api_wait=not bool(original))
        for action in ('validate', 'run'):
            if not spec.get('paths', {}).get(API_PATH + '/' + action, {}).get('post'):
                raise ValueError('Required forecast API endpoint missing')
        for index, (source, row) in enumerate(zip(files, report['files']), 1):
            progress.index, progress.file = index, source.name
            row['status'] = 'running'
            publish_report(args.out, report)
            try:
                records, audit = progress.run('XLSX einlesen', read_workbook, source)
                if audit['sha256'] != row['sha256']:
                    raise ValueError('Source file changed during test')
                dataset = progress.run('Zeitreihe aufbereiten', dataset_from_records, records, audit, args.unit, args.time_basis)
                if config.get('weather_dataset_id'):
                    dataset['weather_region'] = args.weather_region
                row['input'] = audit
                effective = {**config, 'history_from': dataset['period_from'],
                             'forecast_period_from': args.start.isoformat(), 'forecast_period_until': args.until.isoformat()}
                row['effective_configuration'] = effective
                payload = {'datasets': [dataset], 'configuration': effective, 'payload_fit_confirmed': False}
                validation = progress.run('API-Validierung', api.call, f'{index:03d}-validate', API_PATH + '/validate', payload, api_wait=not bool(original))
                row['validation'] = validation
                if validation.get('status') != 'data_validated':
                    raise ValueError('API rejected input validation')
                payload['payload_fit_confirmed'] = True
                def record_job(descriptor):
                    row['job'] = descriptor
                    publish_report(args.out, report)
                response = run_forecast_job(api, f'{index:03d}-forecast', payload, progress, record_job,
                                            args.poll_interval, args.job_timeout)
                metrics, result = progress.run('Ergebnis prüfen und Metriken berechnen', score_response, response, dataset, config, args.start, args.until)
                comparison = metrics.pop('filter_comparison', None)
                if comparison is not None:
                    row['filter_comparison'] = comparison
                row.update(interpret_readiness(metrics, result, args.unit, args.start, args.until))
                row.update(status='completed', metrics=metrics,
                           model_version=result['forecast_run']['model_version'],
                           forecast_run=result['forecast_run'],
                           feature_usage=feature_usage(result),
                           warnings=result.get('readiness_dossier', {}).get('warnings', []))
            except Exception as error:
                row.update(status='failed', error=str(error))
            if row['status'] == 'completed':
                progress.completed += 1
                progress.emit(f"Datei abgeschlossen: RMSE={row['metrics']['rmse']:.6f}, MAE={row['metrics']['mae']:.6f}")
            else:
                progress.failed += 1
                progress.emit('Datei fehlgeschlagen: ' + row.get('error', 'unbekannt'))
            publish_report(args.out, report)
        versions = {r['model_version'] for r in report['files'] if r['status'] == 'completed'}
        if len(versions) > 1:
            raise ValueError('Mixed API model versions in one test run')
        report['status'] = 'completed' if all(r['status'] == 'completed' for r in report['files']) else 'incomplete'
    except KeyboardInterrupt:
        report.update(status='interrupted', error='Interrupted; partial report only')
    except Exception as error:
        report.update(status='failed', error=str(error))
    report['artifacts'] = archive_run_inputs(args.out, script_bytes, config, input_bytes)
    report['finished_at'] = iso(dt.datetime.now(UTC))
    publish_report(args.out, report)
    progress.emit('Lauf beendet: ' + report['status'])
    print(markdown(report), end='')
    return 0 if report['status'] == 'completed' else 1


if __name__ == '__main__':
    sys.exit(main())
