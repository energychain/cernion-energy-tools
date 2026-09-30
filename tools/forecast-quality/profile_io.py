"""Strict XLSX PT15M import; no forecasting code or server dependencies."""
import collections
import datetime as dt
import hashlib
import json
import math
from zoneinfo import ZoneInfo

UTC = dt.timezone.utc
STEP = dt.timedelta(minutes=15)

def iso(value):
    return value.astimezone(UTC).isoformat(timespec='seconds').replace('+00:00', 'Z')


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n')


def excel_datetime(value):
    from openpyxl.utils.datetime import from_excel
    if isinstance(value, (float, int)):
        value = from_excel(value)
    if not isinstance(value, dt.datetime) or value.tzinfo is not None:
        raise ValueError('Expected a naive Excel datetime')
    rounded = value.replace(microsecond=0)
    if value.microsecond >= 999000:
        rounded += dt.timedelta(seconds=1)
    elif value.microsecond > 1000:
        raise ValueError('Timestamp has material fractional seconds')
    if rounded.second or rounded.minute % 15:
        raise ValueError('Timestamp is not aligned to PT15M')
    return rounded


def read_workbook(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Exactly one worksheet expected')
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        columns = {name: headers.index(name) for name in
                   ['Meldepunkt', 'OBIS', 'Datum von', 'Datum bis', 'Wert']}
        records, errors = [], []
        skipped = 0
        meters, obis = set(), set()
        for number, row in enumerate(rows, 2):
            if all(v is None for v in row):
                skipped += 1
                continue
            try:
                data = {k: row[i] if i < len(row) else None for k, i in columns.items()}
                start, end = excel_datetime(data['Datum von']), excel_datetime(data['Datum bis'])
                value = data['Wert']
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                    raise ValueError('Missing/non-numeric Wert; not replaced by zero')
                if data['Meldepunkt'] is None or data['OBIS'] is None:
                    raise ValueError('Missing meter or OBIS')
                meters.add(str(data['Meldepunkt']))
                obis.add(str(data['OBIS']))
                records.append((start, end, float(value), number))
            except (ValueError, TypeError, IndexError) as error:
                if len(errors) < 20:
                    errors.append({'row': number, 'error': str(error)})
        if errors:
            raise ValueError(f'Invalid input rows (first 20): {errors}')
        if not records or len(meters) != 1 or len(obis) != 1:
            raise ValueError('One nonempty meter/OBIS series per workbook required')
        counts = collections.Counter(r[0] for r in records)
        audit = {
            'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'sheet': sheet.title, 'records': len(records), 'blank_rows_ignored': skipped,
            'meter': next(iter(meters)), 'obis': next(iter(obis)),
            'first_wall_time': min(counts).isoformat(), 'last_wall_time': max(counts).isoformat(),
            'duplicate_wall_times': sum(n - 1 for n in counts.values()),
            'non_15_minute_wall_durations': sum(e - s != STEP for s, e, _, _ in records),
        }
        return records, audit
    finally:
        workbook.close()


def local_candidates(wall, zone):
    candidates = set()
    for fold in (0, 1):
        candidate = wall.replace(tzinfo=zone, fold=fold).astimezone(UTC)
        if candidate.astimezone(zone).replace(tzinfo=None) == wall:
            candidates.add(candidate)
    return sorted(candidates)


def dataset_from_records(records, audit, unit, basis):
    zone = ZoneInfo({'berlin': 'Europe/Berlin', 'cet': 'Etc/GMT-1', 'utc': 'UTC'}[basis])
    counts = collections.Counter(r[0] for r in records)
    occurrences = collections.Counter()
    values, seen = [], set()
    previous = None
    for start, end, value, row in records:
        candidates = local_candidates(start, zone)
        if not candidates:
            raise ValueError(f'Row {row}: nonexistent local time {start}; clarify export time convention')
        if len(candidates) == 2:
            if counts[start] != 2:
                raise ValueError(f'Row {row}: ambiguous autumn time {start} without two occurrences')
            instant = candidates[occurrences[start]]
            occurrences[start] += 1
        else:
            instant = candidates[0]
        if instant in seen:
            raise ValueError(f'Row {row}: duplicate physical interval; no deduplication performed')
        if previous is not None and instant <= previous:
            raise ValueError(f'Row {row}: input is not chronologically ordered')
        if instant + STEP not in local_candidates(end, zone):
            raise ValueError(f'Row {row}: Datum bis is not the end of a physical PT15M interval')
        seen.add(instant)
        previous = instant
        values.append({'timestamp': iso(instant), 'value': value})
    audit.update(unit=unit, time_basis=basis, timezone=zone.key,
                 interval_semantics='Datum von=start; Datum bis=exclusive end')
    return {'series_id': 'meter-' + audit['meter'], 'unit': unit, 'timezone': zone.key,
            'period_from': records[0][0].date().isoformat(),
            'period_until': records[-1][0].date().isoformat(), 'values': values}
